"""Bitácora DOCX importer (#7). Synthetic document and names only; every
Claude call is a fake (no network, no API key)."""

import json
import threading
from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import select

from app.models.audit import AuditLog
from app.models.doctor import Doctor
from app.models.log_entry import LogEntry
from app.services.bitacora_import import docx_parse, extract
from app.services.bitacora_import.docx_parse import (
    HeaderHint,
    RawEntry,
    docx_to_markdown_lines,
    header_hint,
    parse_date_line,
    split_entries,
)
from app.services.bitacora_import.extract import SCHEMA, Usage, extract_all, normalize
from app.services.bitacora_import.load import load_staging, read_staging
from scripts import import_bitacora

FAKE_PATIENT_HEADER = "FULANO DE PRUEBA SINTÉTICO CIP 0-000-0000"
NBSP = " "


def _para(doc, *runs):
    """``runs``: plain strings, or (text, {"bold": True, "underline": True})."""
    p = doc.add_paragraph()
    for item in runs:
        text, fmt = (item, {}) if isinstance(item, str) else item
        run = p.add_run(text)
        run.bold = fmt.get("bold")
        run.underline = fmt.get("underline")
    return p


B = {"bold": True}
U = {"underline": True}


@pytest.fixture
def docx_path(tmp_path):
    doc = Document()
    _para(doc, ("BITACORA 2026", B))
    _para(doc, (FAKE_PATIENT_HEADER, B))
    _para(doc, NBSP)
    _para(doc, "22 SEP 2026")
    _para(doc, ("Dra. Ana Prueba", B))
    _para(doc, "Control general. Continuar ", ("Medicamento A", B), " 1 tableta en la mañana.")
    _para(doc, NBSP + " ")
    _para(doc, "Próxima Cita 1400 horas Mar 13 ABR 27.")
    _para(doc, "09 JUL 26")
    _para(doc, ("Dr. Beto Ficticio - WhatsApp", B))
    _para(doc, "Suspender ", ("Medicamento B", U), ". Internet recomienda tomar agua.")
    _para(doc, "21 JUL 2026")  # out of order, as in the real document
    _para(doc, ("FisioCentro - correo", B))
    _para(doc, "Envían plan de ejercicios.")
    _para(doc, "02 FEB 2026")
    _para(doc, "Terapia respiratoria")
    _para(doc, "Sesión de nebulización.")
    _para(doc, "15 ENE 2026")
    _para(doc, ("Resfriado", B), ". Se aplica crema en la zona.")
    _para(doc, ("BITACORA 2021", B))
    _para(doc, FAKE_PATIENT_HEADER)
    _para(doc, "17 DIC 21")
    _para(doc, "Dr. Zeta Inventado: Continuar Medicamento C. Próxima cita 2:00 PM 18 MAY 2022.")
    _para(doc, "03 nov 2021")
    _para(doc, ("Dr. Zeta Inventado:", B))
    _para(doc, "Revisión de exámenes.")
    path = tmp_path / "synthetic.docx"
    doc.save(path)
    return path


# ── Parsing and splitting ────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("22 SEP 2026", date(2026, 9, 22)),
        ("09 JUL 26", date(2026, 7, 9)),
        ("21 JUL 2026", date(2026, 7, 21)),
        ("17 DIC 21", date(2021, 12, 17)),
        ("**3 ago 2024**", date(2024, 8, 3)),
        ("5 SET. 2023", date(2023, 9, 5)),
        (f"1{NBSP}ENE{NBSP}2025", date(2025, 1, 1)),
        ("31 FEB 2026", None),
        ("22 SEP 2026 control", None),
        ("Próxima cita 13 ABR 27", None),
        ("22 XYZ 2026", None),
    ],
)
def test_parse_date_line(line, expected):
    assert parse_date_line(line) == expected


def test_docx_split(docx_path):
    lines = docx_to_markdown_lines(docx_path)
    split = split_entries(lines)
    assert split.years == [2026, 2021]
    assert split.skipped_lines == 2  # the two patient headers
    assert [e.entry_date for e in split.entries] == [
        date(2026, 9, 22),
        date(2026, 7, 9),
        date(2026, 7, 21),
        date(2026, 2, 2),
        date(2026, 1, 15),
        date(2021, 12, 17),
        date(2021, 11, 3),
    ]
    assert [e.index for e in split.entries] == list(range(7))
    first = split.entries[0]
    # Bold kept; paragraphs separated by one blank line (NBSP-only paragraph
    # adds nothing); header dropped.
    assert first.lines == [
        "**Dra. Ana Prueba**",
        "",
        "Control general. Continuar **Medicamento A** 1 tableta en la mañana.",
        "",
        "Próxima Cita 1400 horas Mar 13 ABR 27.",
    ]
    assert "<u>Medicamento B</u>" in split.entries[1].markdown
    for entry in split.entries:
        assert "PRUEBA SINTÉTICO" not in entry.markdown
        assert "0-000-0000" not in entry.markdown
    assert split.entries[0].section_year == 2026
    assert split.entries[5].section_year == 2021
    # Hash is stable and depends on the text.
    again = split_entries(docx_to_markdown_lines(docx_path))
    assert [e.text_sha256 for e in again.entries] == [e.text_sha256 for e in split.entries]
    assert len({e.text_sha256 for e in split.entries}) == 7


def test_split_flags_year_mismatch_and_empty():
    split = split_entries(["BITACORA 2025", "header", "01 ENE 2024", "texto", "02 ENE 2025"])
    assert split.entries[0].warnings == ["date_outside_year_section"]
    assert split.entries[1].warnings == ["empty_entry"]


def test_header_hints(docx_path):
    entries = split_entries(docx_to_markdown_lines(docx_path)).entries
    hints = [header_hint(e.lines) for e in entries]
    assert hints[0] == HeaderHint("Dra. Ana Prueba", True, None)
    assert hints[1] == HeaderHint("Dr. Beto Ficticio", True, "whatsapp")
    assert hints[2] == HeaderHint("FisioCentro", False, "correo")
    assert hints[3] == HeaderHint("Terapia respiratoria", False, None)
    assert hints[4] == HeaderHint()  # no doctor: the title comes directly
    assert hints[5] == HeaderHint("Dr. Zeta Inventado", True, None)  # same-line text
    assert hints[6] == HeaderHint("Dr. Zeta Inventado", True, None)  # trailing colon


@pytest.mark.parametrize(
    ("line", "channel"),
    [
        ("**Dra. X - Teléfono**", "telefono"),
        ("**Dr. Y – correo electrónico**", "correo"),
        ("**Dr. Y (WhatsApp)**", "whatsapp"),
        ("**Dr. Y - visita domiciliaria**", "domicilio"),
        ("**Dr. Y**", None),
    ],
)
def test_channel_suffixes(line, channel):
    assert header_hint([line]).channel == channel


# ── Extraction (fake client) ─────────────────────────────────────────────────


def _model_output(**overrides):
    data = {
        "entry_date": "2026-09-22",
        "doctor": "Dra. Ana Prueba",
        "source_label": "Dra. Ana Prueba",
        "entry_type": "consulta_control",
        "channel": "presencial",
        "title": "Control general",
        "indications": "Continuar Medicamento A 1 tableta en la mañana.",
        "family_notes": None,
        "next_appointment": {
            "date": "2027-04-13",
            "time": "14:00",
            "doctor": "Dra. Ana Prueba",
            "text": "Próxima Cita 1400 horas Mar 13 ABR 27.",
        },
        "medication_events": [],
    }
    data.update(overrides)
    return data


class FakeMessages:
    def __init__(self, respond):
        self.respond = respond
        self.calls = []
        self.lock = threading.Lock()
        self.in_flight = 0
        self.max_in_flight = 0

    def create(self, **params):
        with self.lock:
            self.calls.append(params)
            self.in_flight += 1
            self.max_in_flight = max(self.max_in_flight, self.in_flight)
        try:
            content = params["messages"][0]["content"]
            result = self.respond(content)
            if isinstance(result, Exception):
                raise result
            stop_reason, payload = result
            return SimpleNamespace(
                stop_reason=stop_reason,
                content=[SimpleNamespace(type="text", text=payload)],
                usage=SimpleNamespace(
                    input_tokens=100,
                    output_tokens=50,
                    cache_creation_input_tokens=0,
                    cache_read_input_tokens=1000,
                ),
            )
        finally:
            with self.lock:
                self.in_flight -= 1


def fake_client(respond):
    messages = FakeMessages(respond)
    return SimpleNamespace(beta=SimpleNamespace(messages=messages)), messages


def _respond_by_date(outputs):
    """Answer with the output registered for the entry's date line."""

    def respond(content):
        for date_line, output in outputs.items():
            if f"Línea de fecha: {date_line} " in content:
                return output
        return ("end_turn", json.dumps(_model_output(doctor=None, title="Otro")))

    return respond


def _raw(lines, day=date(2026, 9, 22)):
    return RawEntry(
        index=0, section_year=2026, date_line="22 SEP 2026", entry_date=day, lines=lines
    )


def test_request_params():
    raw = _raw(["**Dra. Ana Prueba**", "Texto"])
    params = extract.request_params("claude-opus-5-5", raw, header_hint(raw.lines))
    assert params["model"] == "claude-opus-5-5"
    assert params["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert params["output_config"]["format"] == {"type": "json_schema", "schema": SCHEMA}
    assert params["fallbacks"] == "default"
    content = params["messages"][0]["content"]
    assert "entry_date = 2026-09-22" in content and "**Dra. Ana Prueba**" in content

    # The schema is strict-compatible: every object closed, every property required.
    def walk(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node["additionalProperties"] is False
                assert set(node["required"]) == set(node["properties"])
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)

    walk(SCHEMA)
    assert "fallbacks" not in extract.request_params("claude-haiku-4-5", raw, HeaderHint())


def test_normalize_rules():
    raw = _raw(["**Dr. Beto Ficticio - WhatsApp**", "x"])
    hint = header_hint(raw.lines)
    out, warnings = normalize(
        _model_output(
            entry_date="2026-09-23",
            channel="presencial",
            entry_type="desconocido",
            title="t" * 400,
            indications="  ",
            next_appointment={"date": "2027-13-40", "time": "25:00", "doctor": None, "text": "x"},
            medication_events=[
                {"kind": "stop", "medication": "Medicamento B", "detail": None},
                {"kind": "bogus", "medication": "Z", "detail": None},
            ],
        ),
        raw,
        hint,
    )
    assert out["entry_date"] == "2026-09-22"  # the splitter's date wins
    assert out["channel"] == "whatsapp"  # explicit header suffix wins
    assert out["entry_type"] == "otro"
    assert len(out["title"]) == 300
    assert out["indications"] is None
    assert out["next_appointment"] == {"date": None, "time": None, "doctor": None, "text": "x"}
    assert out["medication_events"] == [
        {"kind": "stop", "medication": "Medicamento B", "detail": None}
    ]
    assert set(warnings) >= {
        "model_date_differs",
        "channel_from_header",
        "entry_type_unknown",
        "title_truncated",
        "next_appointment_bad_date",
        "next_appointment_bad_time",
    }


def test_normalize_falls_back_to_header_line():
    raw = _raw(["Terapia respiratoria", "Sesión."])
    out, _ = normalize(
        _model_output(doctor=None, source_label=None, title=None), raw, header_hint(raw.lines)
    )
    assert out["source_label"] == "Terapia respiratoria"


def test_extract_all_errors_and_concurrency(docx_path):
    import anthropic
    import httpx2

    entries = split_entries(docx_to_markdown_lines(docx_path)).entries
    items = [(e, header_hint(e.lines)) for e in entries]
    request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
    client, messages = fake_client(
        _respond_by_date(
            {
                "09 JUL 26": ("refusal", ""),
                "21 JUL 2026": ("end_turn", "not json"),
                "02 FEB 2026": anthropic.APIConnectionError(request=request),
                "15 ENE 2026": ("max_tokens", "{"),
                "17 DIC 21": RuntimeError("boom"),
            }
        )
    )
    usage = Usage()
    results = extract_all(client, "claude-opus-5-5", items, usage=usage, concurrency=2)
    assert {i: r.error for i, r in results.items()} == {
        0: None,
        1: "refusal",
        2: "invalid_json",
        3: "api_connection_error",
        4: "truncated",
        5: "RuntimeError",
        6: None,
    }
    assert len(messages.calls) == 7
    assert messages.max_in_flight <= 2
    assert usage.requests == 5  # the two exceptions return no usage
    assert usage.cost_usd("claude-opus-5-5") == pytest.approx(
        (500 * 4 + 250 * 20 + 5000 * 0.20) / 1_000_000
    )
    assert usage.cost_usd("some-other-model") is None


# ── CLI: dry run, staging, load, idempotency ─────────────────────────────────


def _outputs():
    return {
        "22 SEP 2026": ("end_turn", json.dumps(_model_output())),
        "09 JUL 26": (
            "end_turn",
            json.dumps(
                _model_output(
                    entry_date="2026-07-09",
                    doctor="Dr. Beto Ficticio",
                    source_label="Dr. Beto Ficticio - WhatsApp",
                    entry_type="teleconsulta",
                    channel=None,
                    title="Suspender medicamento",
                    indications="Suspender Medicamento B.",
                    family_notes="Internet recomienda tomar agua.",
                    next_appointment=None,
                    medication_events=[
                        {"kind": "stop", "medication": "Medicamento B", "detail": None}
                    ],
                )
            ),
        ),
        "21 JUL 2026": (
            "end_turn",
            json.dumps(
                _model_output(
                    entry_date="2026-07-21",
                    doctor=None,
                    source_label="FisioCentro - correo",
                    entry_type="terapia",
                    channel="correo",
                    title="Plan de ejercicios",
                    next_appointment=None,
                )
            ),
        ),
        "17 DIC 21": (
            "end_turn",
            json.dumps(
                _model_output(
                    entry_date="2021-12-17",
                    doctor="Dr. Zeta Inventado",
                    source_label="Dr. Zeta Inventado",
                    title=None,
                    next_appointment={
                        "date": "2022-05-18",
                        "time": "14:00",
                        "doctor": None,
                        "text": "Próxima cita 2:00 PM 18 MAY 2022.",
                    },
                )
            ),
        ),
        "03 nov 2021": (
            "end_turn",
            json.dumps(
                _model_output(
                    entry_date="2021-11-03",
                    doctor="Zeta Inventado",  # same doctor, written without title
                    source_label="Dr. Zeta Inventado",
                    entry_type="examen",
                    next_appointment=None,
                )
            ),
        ),
        "15 ENE 2026": ("refusal", ""),
    }


@pytest.fixture
def cli(monkeypatch, test_engine):
    from sqlalchemy.orm import sessionmaker

    from app.core.config import Settings

    settings = Settings(anthropic_api_key="test-key", extraction_model="claude-opus-5-5")
    monkeypatch.setattr(import_bitacora, "get_settings", lambda: settings)
    monkeypatch.setattr(import_bitacora, "SessionLocal", sessionmaker(bind=test_engine))
    client, messages = fake_client(_respond_by_date(_outputs()))
    monkeypatch.setattr(import_bitacora, "make_client", lambda key: client)
    return SimpleNamespace(settings=settings, messages=messages)


def test_dry_run_writes_staging_and_touches_no_db(cli, docx_path, tmp_path, monkeypatch, capsys):
    def no_db():
        raise AssertionError("dry run must not open the database")

    monkeypatch.setattr(import_bitacora, "SessionLocal", no_db)
    staging_path = tmp_path / "private" / "staging.json"
    code = import_bitacora.main(
        [str(docx_path), "--patient", "papa", "--dry-run", "--staging", str(staging_path)]
    )
    assert code == 0
    assert len(cli.messages.calls) == 7
    staging = read_staging(staging_path)
    assert (staging_path.stat().st_mode & 0o777) == 0o600
    assert staging["patient"] == "papa"
    assert staging["years"] == [2026, 2021]
    assert [e["error"] for e in staging["entries"]].count("refusal") == 1
    raw_text = staging_path.read_text()
    assert "PRUEBA SINTÉTICO" not in raw_text and "0-000-0000" not in raw_text
    second = staging["entries"][1]
    assert second["extraction"]["channel"] == "whatsapp"
    assert second["header_hint"]["channel"] == "whatsapp"
    out = capsys.readouterr()
    assert "estimated cost: US$" in out.out
    assert "PATIENT HEALTH INFORMATION" in out.err
    # No entry text in the console output.
    assert "Medicamento" not in out.out + out.err


def test_limit_and_default_staging_path(cli, docx_path, tmp_path, monkeypatch):
    (tmp_path / ".git").mkdir()  # a checkout: private/ is gitignored and persists
    monkeypatch.chdir(tmp_path)
    assert (
        import_bitacora.main([str(docx_path), "--patient", "papa", "--dry-run", "--limit", "2"])
        == 0
    )
    assert len(cli.messages.calls) == 2
    [path] = (tmp_path / "private").glob("bitacora-staging-*.json")
    assert len(read_staging(path)["entries"]) == 2


def test_missing_api_key(cli, docx_path, tmp_path):
    cli.settings.anthropic_api_key = ""
    args = [str(docx_path), "--patient", "papa", "--dry-run", "--staging", str(tmp_path / "s.json")]
    assert import_bitacora.main(args) == 1
    assert cli.messages.calls == []


def test_from_staging_load_and_idempotency(
    cli, test_db, patients, member, docx_path, tmp_path, capsys
):
    ambiguous = [
        Doctor(name="Dra. Ana Prueba"),
        Doctor(name="Dra. Ana María Prueba", aliases=["Ana Prueba"]),
    ]
    test_db.add_all(ambiguous)
    test_db.commit()
    staging_path = tmp_path / "staging.json"
    base = [str(docx_path), "--patient", "papa"]
    assert import_bitacora.main([*base, "--dry-run", "--staging", str(staging_path)]) == 0
    capsys.readouterr()

    load = ["--patient", "papa", "--from-staging", str(staging_path), "--by", member.email]
    assert import_bitacora.main(load) == 2  # the refused entry fails
    out = capsys.readouterr().out
    assert "loaded: 6 new, 0 skipped, 1 failed" in out
    assert "failed entry #4: refusal" in out
    assert "Dr. Beto Ficticio" in out and "Dr. Zeta Inventado" in out
    assert "Dra. Ana Prueba" in out  # ambiguous name reported

    rows = test_db.scalars(select(LogEntry).order_by(LogEntry.id)).all()
    assert len(rows) == 6
    assert {r.review_status for r in rows} == {"pending_review"}
    assert {r.origin for r in rows} == {"import_docx"}
    assert {r.created_by for r in rows} == {member.id}
    assert {r.patient_id for r in rows} == {patients[0].id}

    by_index = {r.origin_ref["entry_index"]: r for r in rows}
    first = by_index[0]
    assert first.doctor_id is None  # ambiguous: no doctor, name kept in the label
    assert first.source_label == "Dra. Ana Prueba"
    assert first.entry_date == date(2026, 9, 22)
    assert first.origin_ref["next_appointment"]["date"] == "2027-04-13"
    assert first.origin_ref["original_markdown"].startswith("**Dra. Ana Prueba**")
    assert set(first.origin_ref) >= {"source_sha256", "text_sha256", "medication_events"}

    beto = by_index[1]
    assert beto.doctor.name == "Dr. Beto Ficticio"
    assert beto.channel == "whatsapp"
    assert beto.family_notes == "Internet recomienda tomar agua."
    assert beto.origin_ref["medication_events"][0]["kind"] == "stop"

    rehab = by_index[2]
    assert (rehab.doctor_id, rehab.source_label, rehab.channel) == (
        None,
        "FisioCentro - correo",
        "correo",
    )
    # "Dr. Zeta Inventado" created once and matched again for "Zeta Inventado".
    assert by_index[5].doctor_id == by_index[6].doctor_id is not None
    created = test_db.scalars(select(Doctor).where(Doctor.id.notin_([d.id for d in ambiguous])))
    assert sorted(d.name for d in created) == ["Dr. Beto Ficticio", "Dr. Zeta Inventado"]
    audit_doctors = test_db.scalars(select(AuditLog).where(AuditLog.entity == "doctor")).all()
    assert len(audit_doctors) == 2

    # Re-running the load skips everything already imported.
    assert import_bitacora.main(load) == 2
    assert "loaded: 0 new, 6 skipped, 1 failed" in capsys.readouterr().out
    assert test_db.scalar(select(LogEntry.id).order_by(LogEntry.id.desc())) == rows[-1].id

    # A full run does not send already-imported entries to Claude again.
    calls_before = len(cli.messages.calls)
    again = [*base, "--staging", str(tmp_path / "again.json"), "--by", member.email]
    assert import_bitacora.main(again) == 2
    assert len(cli.messages.calls) - calls_before == 1  # only the refused one
    assert "already imported: 6 entries" in capsys.readouterr().out


def test_full_run_and_mismatches(cli, test_db, patients, member, docx_path, tmp_path, capsys):
    staging_path = tmp_path / "s.json"
    by = ["--by", member.email]
    run = [str(docx_path), "--patient", "papa", "--staging", str(staging_path)]
    assert import_bitacora.main(run) == 1  # --by is required to load
    assert cli.messages.calls == []
    assert "--by" in capsys.readouterr().err
    assert import_bitacora.main([*run, *by]) == 2
    assert "loaded: 6 new, 0 skipped, 1 failed" in capsys.readouterr().out
    assert test_db.scalar(select(LogEntry.id).limit(1)) is not None
    # Staging for another patient, or another .docx, is refused.
    assert (
        import_bitacora.main(["--patient", "mama", "--from-staging", str(staging_path), *by]) == 1
    )
    assert "not 'mama'" in capsys.readouterr().err
    other = tmp_path / "other.docx"
    Document().save(other)
    assert (
        import_bitacora.main(
            [str(other), "--patient", "papa", "--from-staging", str(staging_path), *by]
        )
        == 1
    )
    assert "does not match" in capsys.readouterr().err
    assert import_bitacora.main(["--patient", "papa", *by]) == 1
    missing = ["--patient", "papa", "--from-staging", str(tmp_path / "x"), *by]
    assert import_bitacora.main(missing) == 1
    assert import_bitacora.main(["--patient", "papa", "--from-staging", str(staging_path)]) == 1


def test_invalid_entry_does_not_abort(test_db, patients):
    staging = {
        "version": 1,
        "source_sha256": "a" * 64,
        "patient": "papa",
        "model": "m",
        "entries": [
            {
                "entry_index": 0,
                "text_sha256": "b",
                "date_line": "01 ENE 2099",
                "entry_date": "2099-01-01",  # future: rejected by validate()
                "markdown": "x",
                "extraction": _model_output(),
                "error": None,
                "warnings": [],
            },
            {
                "entry_index": 1,
                "text_sha256": "c",
                "date_line": "01 ENE 2026",
                "entry_date": "2026-01-01",
                "markdown": "y",
                "extraction": _model_output(doctor=None, source_label=None, title="Nota"),
                "error": None,
                "warnings": [],
            },
        ],
    }
    report = load_staging(test_db, staging, patient=patients[0], user_id=None)
    assert (report.new, report.failed) == (1, 1)
    assert "hechos pasados" in report.failures[0][1]
    assert test_db.scalar(select(LogEntry.title)) == "Nota"


def test_docx_module_never_logs(docx_path, caplog):
    split_entries(docx_to_markdown_lines(docx_path))
    assert caplog.records == []
    assert docx_parse.MONTHS["DIC"] == 12


# ── Hardening: PHI guard, markdown fields, staging safety, resume ───────────


def test_loose_section_headers_never_become_entry_text():
    split = split_entries(
        [
            "BITACORA 2026",
            FAKE_PATIENT_HEADER,
            "22 SEP 2026",
            "Texto",
            "",
            "**BITACORA 2025 (cont.)**",
            FAKE_PATIENT_HEADER,
            "01 MAR 2025",
            "Otro texto",
            "Bitácora 2024 FULANO DE PRUEBA",
            "Línea suelta del encabezado",
            "05 ENE 2024",
            "Más texto",
            "BITACORA",
            "07 ENE 2023",
            "Sin año",
        ]
    )
    assert split.years == [2026, 2025, 2024, None]
    assert split.counts_by_section() == [(2026, 1), (2025, 1), (2024, 1), (None, 1)]
    assert split.skipped_lines == 3
    assert [e.lines for e in split.entries] == [
        ["Texto"],
        ["Otro texto"],
        ["Más texto"],
        ["Sin año"],
    ]
    assert split.leaked_entries == []


def test_repeated_patient_header_inside_entry_is_flagged():
    split = split_entries(
        [
            "BITACORA 2026",
            FAKE_PATIENT_HEADER,
            "22 SEP 2026",
            "Texto",
            "21 SEP 2026",
            f"**{FAKE_PATIENT_HEADER.lower()}**",  # pasted again inside an entry
        ]
    )
    assert split.leaked_entries == [1]


def test_cli_aborts_on_leaked_header_before_any_call(cli, tmp_path, capsys):
    doc = Document()
    for line in ["BITACORA 2026", FAKE_PATIENT_HEADER, "22 SEP 2026", FAKE_PATIENT_HEADER]:
        _para(doc, line)
    path = tmp_path / "leak.docx"
    doc.save(path)
    staging = tmp_path / "s.json"
    args = [str(path), "--patient", "papa", "--dry-run", "--staging", str(staging)]
    assert import_bitacora.main(args) == 1
    assert cli.messages.calls == []
    assert not staging.exists()
    err = capsys.readouterr().err
    assert "reappears verbatim inside entries 0" in err
    assert "PRUEBA" not in err


def test_markdown_fields_keep_bold_strip_underline():
    raw = _raw(["**Dra. Ana Prueba**", "", "texto"])
    out, _ = normalize(
        _model_output(
            doctor="**Dra. Ana Prueba**",
            title="<u>**Control**</u>",
            indications="Continuar **Medicamento A**.\n\n<u>Dieta</u> blanda.",
            family_notes="<U>Nota</U> de la familia.",
        ),
        raw,
        header_hint(raw.lines),
    )
    assert out["doctor"] == "Dra. Ana Prueba"
    assert out["title"] == "Control"
    assert out["indications"] == "Continuar **Medicamento A**.\n\nDieta blanda."
    assert out["family_notes"] == "Nota de la familia."
    assert "línea en blanco" in extract.SYSTEM_PROMPT


def _one_entry_staging(**extraction):
    return {
        "version": 1,
        "source_sha256": "d" * 64,
        "patient": "papa",
        "model": "m",
        "entries": [
            {
                "entry_index": 0,
                "text_sha256": "e",
                "date_line": "01 ENE 2026",
                "entry_date": "2026-01-01",
                "markdown": "x",
                "extraction": _model_output(**extraction),
                "error": None,
                "warnings": [],
            }
        ],
    }


def test_title_only_doctor_is_never_created(test_db, patients):
    staging = _one_entry_staging(doctor="Dra.", source_label=None, title=None)
    report = load_staging(test_db, staging, patient=patients[0], user_id=None)
    assert report.new == 1 and report.created_doctors == []
    assert test_db.scalar(select(Doctor.id)) is None
    entry = test_db.scalar(select(LogEntry))
    assert (entry.doctor_id, entry.source_label) == (None, "Dra.")


def test_other_source_duplicates_are_reported(test_db, patients):
    first = _one_entry_staging()
    assert load_staging(test_db, first, patient=patients[0], user_id=None).new == 1
    resaved = {**first, "source_sha256": "f" * 64}
    report = load_staging(test_db, resaved, patient=patients[0], user_id=None)
    assert report.new == 1 and report.other_source_duplicates == [0]


def test_staging_mode_is_0600_when_overwriting(tmp_path):
    from app.services.bitacora_import.load import write_staging

    path = tmp_path / "s.json"
    path.write_text("{}")
    path.chmod(0o644)
    write_staging(path, {"version": 1})
    assert (path.stat().st_mode & 0o777) == 0o600
    assert json.loads(path.read_text()) == {"version": 1}
    assert [p.name for p in tmp_path.iterdir()] == ["s.json"]


def test_staging_required_outside_checkout(cli, docx_path, tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(import_bitacora, "_in_checkout", lambda cwd: False)
    assert import_bitacora.main([str(docx_path), "--patient", "papa", "--dry-run"]) == 1
    assert "--staging is required" in capsys.readouterr().err
    assert cli.messages.calls == []


def test_unwritable_staging_fails_before_calls(cli, docx_path, tmp_path, capsys):
    blocker = tmp_path / "file"
    blocker.write_text("")
    args = [str(docx_path), "--patient", "papa", "--dry-run", "--staging", str(blocker / "s.json")]
    assert import_bitacora.main(args) == 1
    assert "cannot write the staging file" in capsys.readouterr().err
    assert cli.messages.calls == []


def test_resume_from_partial_staging(cli, docx_path, tmp_path, capsys):
    staging = tmp_path / "s.json"
    base = [str(docx_path), "--patient", "papa", "--dry-run", "--staging", str(staging)]
    assert import_bitacora.main([*base, "--limit", "3"]) == 0
    assert len(cli.messages.calls) == 3
    assert import_bitacora.main(base) == 0
    assert len(cli.messages.calls) == 7  # only the 4 not yet extracted
    out = capsys.readouterr().out
    assert "resuming" in out and "3 entries already extracted" in out
    assert "entries per section: 2026: 5, 2021: 2" in out
    assert len(read_staging(staging)["entries"]) == 7
    # The refused entry is retried on the next run; the rest are reused.
    assert import_bitacora.main(base) == 0
    assert len(cli.messages.calls) == 8
    # A staging file for another patient is not overwritten.
    other = [str(docx_path), "--patient", "mama", "--dry-run", "--staging", str(staging)]
    assert import_bitacora.main(other) == 1
    assert len(cli.messages.calls) == 8


def test_partial_staging_is_saved_during_extraction(cli, docx_path, tmp_path, monkeypatch):
    staging = tmp_path / "s.json"
    saves = []
    real_write = import_bitacora.write_staging

    def spy(path, data):
        saves.append(sum(1 for e in data["entries"] if e["error"] != "not_extracted"))
        real_write(path, data)

    monkeypatch.setattr(import_bitacora, "write_staging", spy)
    monkeypatch.setattr(import_bitacora, "SAVE_EVERY", 2)
    args = [str(docx_path), "--patient", "papa", "--dry-run", "--staging", str(staging)]
    assert import_bitacora.main(args) == 0
    assert saves[0] == 0  # written before the first call
    assert saves[1:-1] == [2, 4, 6]
    assert saves[-1] == 7
