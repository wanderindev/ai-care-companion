"""Medication-list DOCX importer (#24). Synthetic document and names only; the
Claude call is a fake (no network, no API key)."""

import json
import stat
from datetime import date, time
from types import SimpleNamespace

import pytest
from docx import Document
from sqlalchemy import select

from app.models.doctor import Doctor
from app.models.medication import MedAvoid, MedItem, MedListVersion
from app.services import meds
from app.services.meds_import import docx_parse, extract
from app.services.meds_import.docx_parse import (
    classify_header,
    docx_to_markdown_lines,
    parse_since,
    split_document,
)
from app.services.meds_import.load import CHANGE_SUMMARY, load_staging, read_staging
from app.services.meds_import.review import review_list
from scripts import import_meds

FAKE_PATIENT_HEADER = "FULANO DE PRUEBA SINTÉTICO CIP 0-000-0000"
B = {"bold": True}


def _para(doc, *runs):
    p = doc.add_paragraph()
    for item in runs:
        text, fmt = (item, {}) if isinstance(item, str) else item
        run = p.add_run(text)
        run.bold = fmt.get("bold")
    return p


def build_docx(path, *, extra=()):
    doc = Document()
    _para(doc, ("LISTA DE MEDICAMENTOS", B))
    _para(doc, (FAKE_PATIENT_HEADER, B))
    _para(doc, "DESDE 30 SEP 26")
    _para(doc, ("AL DESPERTAR", B))
    _para(doc, "Medicina A 250 mg ½ tableta 30 min antes del desayuno")
    _para(doc, ("DESAYUNO", B))
    _para(doc, "Medicina B 1 tableta lunes, miércoles y viernes")
    _para(doc, "Medicina C 1 tableta: vie 25 sep 1; lun 28 sep ½; jue 1 oct ¼")
    _para(doc, ("ENTRE 7:30 y 8:00 PM", B))
    _para(doc, "Medicina D 1 cápsula")
    _para(doc, ("CENA:", B))
    _para(doc, "Rescate: Medicina E 1 tableta si hay dolor")
    _para(doc, (FAKE_PATIENT_HEADER, B))  # page-2 header repeated: dropped
    _para(doc, ("OTROS MEDICAMENTOS Y SUPLEMENTOS", B))
    _para(doc, "Suplemento F 1 sobre")
    for line in extra:
        _para(doc, line)
    _para(doc, ("EVITAR", B))
    _para(doc, "Sustancia X (reacción sintética)")
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Hora", "Medicina A"), ("6:00", "1"), ("12:00", "½")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    doc.save(path)
    return path


@pytest.fixture
def docx_path(tmp_path):
    return build_docx(tmp_path / "lista.docx")


# ── Parsing ──────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("text", "slot", "at", "window"),
    [
        ("AL DESPERTAR", "al_despertar", None, None),
        ("**Cuando despierte:**", "al_despertar", None, None),
        ("Desayuno", "desayuno", None, None),
        ("**MAÑANA**", "manana", None, None),
        ("<u>ALMUERZO</u>", "almuerzo", None, None),
        ("TARDE", "tarde", None, None),
        ("CENA:", "cena", None, None),
        ("NOCHE", "noche", None, None),
        ("Madrugada", "madrugada", None, None),
        ("ENTRE 7:30 y 8:00 PM", "noche", time(19, 30), "Entre 7:30 y 8:00 PM"),
        ("Entre 7:30 y 8:00 p.m.", "noche", time(19, 30), "Entre 7:30 y 8:00 PM"),
        ("6:00 AM", "manana", time(6, 0), "6:00 AM"),
        ("A LAS 3 PM", "tarde", time(15, 0), "A las 3 PM"),
        ("12:30 AM", "madrugada", time(0, 30), "12:30 AM"),
        ("DESAYUNO (7:00 AM)", "desayuno", time(7, 0), "7:00 AM"),
        ("Rescate", "segun_necesidad", None, None),
        ("Según necesidad", "segun_necesidad", None, None),
        ("OTROS MEDICAMENTOS Y SUPLEMENTOS", "otros", None, None),
        ("Otros medicamentos", "otros", None, None),
        ("EVITAR:", docx_parse.AVOID, None, None),
    ],
)
def test_header_to_slot(text, slot, at, window):
    header = classify_header(text)
    assert header is not None
    assert (header.slot, header.time, header.window) == (slot, at, window)


@pytest.mark.parametrize(
    "text", ["Medicina A 1 tableta", "1", "½", "25:00 PM", "INDICACIONES", "", None]
)
def test_not_a_header(text):
    assert classify_header(text) is None


@pytest.mark.parametrize(
    ("rest", "expected"),
    [
        ("30 SEP 26", date(2026, 9, 30)),
        ("30 de septiembre de 2026", date(2026, 9, 30)),
        ("30/09/2026", date(2026, 9, 30)),
        ("el 1 oct 2026", date(2026, 10, 1)),
        ("ayer", None),
        ("31/02/2026", None),
    ],
)
def test_parse_since(rest, expected):
    assert parse_since(rest) == expected


def test_docx_sections_and_phi_guard(docx_path):
    parsed = split_document(docx_to_markdown_lines(docx_path))
    assert [(s.header, s.slot, s.lines, s.table_rows) for s in parsed.sections] == [
        ("AL DESPERTAR", "al_despertar", 1, 0),
        ("DESAYUNO", "desayuno", 2, 0),
        ("ENTRE 7:30 y 8:00 PM", "noche", 1, 0),
        ("CENA", "cena", 1, 0),
        ("OTROS MEDICAMENTOS Y SUPLEMENTOS", "otros", 1, 0),
        ("EVITAR", "evitar", 1, 2),  # one line + a table with 2 data rows
    ]
    assert parsed.sections[2].time == time(19, 30)
    assert parsed.effective_date == date(2026, 9, 30)
    assert parsed.skipped_lines == 3  # title, patient header, DESDE line
    assert parsed.repeated_header_lines == 1
    assert parsed.id_lines == []
    assert parsed.warnings == []
    md = parsed.markdown
    assert md.startswith("[S1] **AL DESPERTAR**")
    assert "\n[S3] **ENTRE 7:30 y 8:00 PM**\n" in md
    assert "PRUEBA SINTÉTICO" not in md and "0-000-0000" not in md and "DESDE" not in md
    assert "| Hora | Medicina A |\n| --- | --- |\n| 6:00 | 1 |\n| 12:00 | ½ |" in md
    assert parsed.section_counts()["desayuno"] == 1


def test_id_document_in_body_is_flagged(tmp_path):
    path = build_docx(tmp_path / "x.docx", extra=["Cédula 0-000-0001"])
    parsed = split_document(docx_to_markdown_lines(path))
    assert len(parsed.id_lines) == 1


def test_no_since_date_and_desde_in_preamble_line():
    parsed = split_document(["LISTA", "AL DESPERTAR", "Medicina A"])
    assert parsed.effective_date is None and parsed.warnings == ["no_since_date"]
    parsed = split_document(["LISTA DE MEDICAMENTOS DESDE 30/09/2026", "CENA", "Medicina A"])
    assert parsed.effective_date == date(2026, 9, 30) and parsed.skipped_lines == 1
    # An item's own "desde" stays in the body (it is the item's start date),
    # also as a DESDE line of its own inside a section.
    parsed = split_document(["CENA", "Medicina A desde el 5 oct"])
    assert parsed.lines == ["[S1] CENA", "Medicina A desde el 5 oct"]
    parsed = split_document(["DESDE 30 SEP 26", "CENA", "Medicina A", "DESDE 5 OCT 26"])
    assert parsed.effective_date == date(2026, 9, 30)
    assert parsed.lines == ["[S1] CENA", "Medicina A", "DESDE 5 OCT 26"]
    assert parsed.sections[0].lines == 2
    parsed = split_document(["CENA", "DESDE 5 OCT 26", "Medicina A"])
    assert parsed.effective_date is None and "DESDE 5 OCT 26" in parsed.lines


def test_table_row_header_and_unknown_preamble():
    lines = ["| DESAYUNO |", "| --- |", "| Medicina A |"]
    parsed = split_document(lines)
    assert [s.slot for s in parsed.sections] == ["desayuno"]
    assert split_document(["Sin encabezados", "Medicina A"]).sections == []


@pytest.mark.parametrize(
    ("text", "slot", "at", "note"),
    [
        ("DESPUÉS DEL DESAYUNO", "desayuno", None, "Después del desayuno"),
        ("**Antes del desayuno:**", "desayuno", None, "Antes del desayuno"),
        ("ANTES DE LA CENA", "cena", None, "Antes de la cena"),
        ("Luego del almuerzo", "almuerzo", None, "Luego del almuerzo"),
        ("ANTES DE ACOSTARSE", "noche", None, "Antes de acostarse"),
        ("DESPUÉS DEL DESAYUNO (8:00 AM)", "desayuno", time(8, 0), "Después del desayuno"),
        ("CON EL ALMUERZO", "almuerzo", None, None),  # already in the table
        ("AL ACOSTARSE", "noche", None, None),
    ],
)
def test_qualified_header(text, slot, at, note):
    header = classify_header(text)
    assert header is not None
    assert (header.slot, header.time, header.note) == (slot, at, note)


@pytest.mark.parametrize(
    "text",
    [
        "ANTES DE EVITAR",
        "DESPUÉS DE OTROS",
        "ANTES DE COMER",
        "DESPUÉS DE ACOSTARSE X",
        # An amount reads as an item's instruction: never a new section.
        "30 MIN ANTES DEL ALMUERZO",
    ],
)
def test_qualifier_without_slot_word_is_not_a_header(text):
    assert classify_header(text) is None


def test_qualified_first_header_keeps_its_section():
    lines = [
        "LISTA DE MEDICAMENTOS",
        FAKE_PATIENT_HEADER,
        "DESDE 30 SEP 26",
        "**DESPUÉS DEL DESAYUNO**",
        "Medicina A 1 tableta",
        "Medicina B 2 tabletas",
        "NOCHE",
        "Medicina C 1 tableta",
    ]
    parsed = split_document(lines)
    assert [(s.slot, s.note, s.lines) for s in parsed.sections] == [
        ("desayuno", "Después del desayuno", 2),
        ("noche", None, 1),
    ]
    assert parsed.preamble_unrecognized == []
    assert parsed.skipped_lines == 3


def test_preamble_unrecognized_lines():
    # Title, a name line (letters only), DESDE: all expected, nothing flagged.
    parsed = split_document(
        ["LISTA DE MEDICAMENTOS", "", "FULANA DE PRUEBA", "DESDE 30 SEP 26", "CENA", "Med A"]
    )
    assert parsed.preamble_unrecognized == [] and parsed.skipped_lines == 3
    # A medication line right under DESDE, before any header: flagged by its
    # position among non-empty lines (blank lines not counted).
    parsed = split_document(
        [FAKE_PATIENT_HEADER, "DESDE 30 SEP 26", "", "Medicina A 1 tableta", "CENA", "Med B"]
    )
    assert parsed.preamble_unrecognized == [3]
    # A letters-only line is the name only once, and only before DESDE.
    parsed = split_document(
        ["FULANA DE PRUEBA", "Vitamina C", "DESDE 30 SEP 26", "Magnesio", "CENA", "Med B"]
    )
    assert parsed.preamble_unrecognized == [2, 4]
    # An unknown header before the first recognized one: its items are flagged.
    parsed = split_document(
        [
            "LISTA",
            FAKE_PATIENT_HEADER,
            "DESDE 30 SEP 26",
            "ENCABEZADO RARO",
            "Medicina A 1 tab",
            "CENA",
            "Med B",
        ]
    )
    assert parsed.preamble_unrecognized == [4, 5]
    # A dosing pattern is not an ID number.
    assert split_document(["Medicina A 1-0-1", "CENA", "Med B"]).preamble_unrecognized == [1]


@pytest.mark.parametrize(
    ("lines", "flagged", "since"),
    [
        # An item's own "desde" is not the list's DESDE (nor does it set it).
        (["LISTA", "DESDE 30 SEP 26", "Medicina A 100 mg desde 01/10/2026"], [3], "2026-09-30"),
        (["LISTA", "Medicina A 100 mg desde 01/10/2026"], [2], None),
        (["LISTA", "Medicina A desde 01/10/2026"], [2], None),
        # A date is not an ID number.
        (
            ["LISTA", FAKE_PATIENT_HEADER, "DESDE 30 SEP 26", "Medicina D 1 tab hasta 15-10-2026"],
            [4],
            "2026-09-30",
        ),
        (
            ["LISTA", FAKE_PATIENT_HEADER, "DESDE 30 SEP 26", "Medicina D hasta 15-10-2026"],
            [4],
            "2026-09-30",
        ),
        # The ID line carries the name: no letters-only line counts as the name after it.
        (
            ["LISTA DE MEDICAMENTOS", FAKE_PATIENT_HEADER, "Omeprazol", "DESDE 30 SEP 26"],
            [3],
            "2026-09-30",
        ),
        # A bare ID line (no name) still leaves room for one name line.
        (["CIP 0-000-0000", "FULANA DE PRUEBA", "DESDE 30 SEP 26"], [], "2026-09-30"),
        # One of each: a second DESDE or ID line is flagged; so is an unreadable DESDE.
        (["DESDE 30 SEP 26", "DESDE 5 OCT 26"], [2], "2026-09-30"),
        (["CIP 0-000-0000", "CIP 0-000-0001"], [2], None),
        (["DESDE AYER"], [1], None),
        (["LISTA DE MEDICAMENTOS DESDE 30/09/2026"], [], "2026-09-30"),
    ],
)
def test_preamble_item_shapes_are_flagged(lines, flagged, since):
    parsed = split_document([*lines, "CENA", "Medicina X"])
    assert parsed.preamble_unrecognized == flagged
    assert (parsed.effective_date.isoformat() if parsed.effective_date else None) == since


# ── Extraction (fake client) ─────────────────────────────────────────────────


def _item(name, section, header, **kw):
    data = {k: None for k in extract._ITEM_PROPS}
    data.update(name=name, section=section, header=header, rescue=False, quantity="1")
    data.update(kw)
    return data


def model_output():
    return {
        "items": [
            _item(
                "Medicina A",
                1,
                "AL DESPERTAR",
                strength="250 mg",
                quantity="½",
                instructions="30 min antes del desayuno",
            ),
            _item("Medicina B", 2, "DESAYUNO", days=[5, 1, 3, 3], prescriber="Dra. Ana Prueba"),
            _item(
                "Medicina C",
                2,
                "DESAYUNO",
                schedule_note="vie 25 sep 1; lun 28 sep ½; jue 1 oct ¼",
                prescriber="Dr. Desconocido",
            ),
            _item("Medicina D", 3, "[S3] ENTRE 7:30 y 8:00 PM", quantity="1 cápsula"),
            _item("Medicina E", 4, "CENA", rescue=True, instructions="si hay dolor"),
            _item("Suplemento F", 5, "OTROS MEDICAMENTOS Y SUPLEMENTOS", quantity="1 sobre"),
            _item("Medicina G", 5, "SECCIÓN RARA", days=[1, 2, 3, 4, 5, 6, 7]),
        ],
        "avoid": [{"substance": "Sustancia X", "reason": "reacción sintética"}],
        "list_note": "| Hora | Medicina A |\n| --- | --- |\n| 6:00 | 1 |",
    }


class FakeMessages:
    def __init__(self, payload, stop_reason="end_turn"):
        self.payload = payload
        self.stop_reason = stop_reason
        self.calls = []

    def create(self, **params):
        self.calls.append(params)
        return SimpleNamespace(
            stop_reason=self.stop_reason,
            content=[SimpleNamespace(type="text", text=json.dumps(self.payload))],
            usage=SimpleNamespace(
                input_tokens=2000,
                output_tokens=3000,
                cache_creation_input_tokens=1500,
                cache_read_input_tokens=0,
            ),
        )


def test_request_params_and_message(docx_path):
    parsed = split_document(docx_to_markdown_lines(docx_path))
    params = extract.request_params("claude-opus-5-5", parsed)
    assert params["output_config"]["format"]["schema"] == extract.SCHEMA
    assert params["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert params["fallbacks"] == "default"
    content = params["messages"][0]["content"]
    assert "Fecha de la lista (DESDE): 2026-09-30" in content
    assert '"ENTRE 7:30 y 8:00 PM"' in content
    assert "PRUEBA SINTÉTICO" not in content and "0-000-0000" not in content


def test_normalize_slot_from_section_number(docx_path):
    parsed = split_document(docx_to_markdown_lines(docx_path))
    output = model_output()
    output["items"] += [
        _item("Sustancia Y", 6, "EVITAR", instructions="alergia sintética"),
        _item("Medicina Z", None, None),
        _item("Medicina W", 99, "CENA"),
        _item("  ", 4, "CENA"),
        _item(
            "Medicina H",
            2,
            "CENA",  # the model's label disagrees with the section number
            days=[0, 9],
            time="25:00",
            start_date="2026-10-05",
            end_date="2026-10-01",
        ),
        _item("Medicina Q", 2, "DESAYUNO", quantity="x" * 60, schedule_note="y" * 5000),
    ]
    data, warnings = extract.normalize(output, parsed)
    by_name = {i["name"]: i for i in data["items"]}
    assert {n: i["slot"] for n, i in by_name.items()} == {
        "Medicina A": "al_despertar",
        "Medicina B": "desayuno",
        "Medicina C": "desayuno",
        "Medicina D": "noche",
        "Medicina E": "segun_necesidad",
        "Suplemento F": "otros",
        "Medicina G": "otros",  # section 5 by position, whatever the label
        "Sustancia Y": None,  # from EVITAR: never an active medication silently
        "Medicina Z": None,  # from no section
        "Medicina W": None,  # section number out of range
        "Medicina H": "desayuno",  # position wins over the label
        "Medicina Q": "desayuno",
    }
    assert by_name["Medicina B"]["days"] == [1, 3, 5]
    assert by_name["Medicina G"]["days"] is None  # all seven days = daily
    assert by_name["Medicina D"]["time"] == "19:30"
    assert by_name["Medicina D"]["instructions"] == "Entre 7:30 y 8:00 PM"
    assert by_name["Medicina E"]["time"] is None
    assert by_name["Medicina H"]["days"] is None and by_name["Medicina H"]["time"] is None
    assert by_name["Medicina H"]["end_date"] == "2026-10-01"  # kept; the load refuses it
    # Never truncated: flagged, kept whole for the load's validation to refuse.
    assert by_name["Medicina Q"]["quantity"] == "x" * 60
    assert by_name["Medicina Q"]["schedule_note"] == "y" * 5000
    assert data["avoid"] == [{"substance": "Sustancia X", "reason": "reacción sintética"}]
    assert data["effective_date"] == "2026-09-30"
    assert data["list_note"].startswith("| Hora |")
    assert "item 10: bad_section" in warnings
    assert "model item 11: no_name (dropped)" in warnings
    expected = {"item 11: bad_days", "item 11: bad_time", "item 11: end_before_start"}
    assert expected <= set(warnings)
    assert "item 12: quantity_too_long" in warnings
    assert not any("schedule_note" in w for w in warnings)  # up to 10000 allowed
    assert not any("Medicina" in w or "Sustancia" in w for w in warnings)

    sections = [{"header": x.header, "slot": x.slot, "lines": x.lines} for x in parsed.sections]
    review = review_list(sections, data, date(2026, 10, 9))
    assert review.blocking == [
        f"item {n}: needs_review (from EVITAR or from no section; set its slot or move it to avoid)"
        for n in (8, 9, 10)
    ]
    assert "item 7: unknown_header" in review.warnings
    assert "item 11: header_mismatch section=2" in review.warnings
    assert "item 4: unknown_header" not in review.warnings  # "[S3] " prefix is ignored
    assert "item 9: no_section" not in review.warnings  # already blocking
    # Unreadable values the extraction dropped, and a past end date, are gated.
    assert "item 11: bad_days (fix the field, then delete the code)" in review.warnings
    assert "item 11: bad_time (fix the field, then delete the code)" in review.warnings
    assert "item 10: bad_section (fix the field, then delete the code)" in review.warnings
    assert "item 11: ends_in_past (would not be on the current list)" in review.warnings
    assert review.notes == []


@pytest.mark.parametrize("rows", [1, 2])
def test_section_header_in_table_header_row(rows):
    # The table's merged title row is itself the section header: it was never
    # counted, so the separator after it must not uncount a data row.
    lines = [
        "| DESAYUNO | DESAYUNO |",
        "| --- | --- |",
        *[f"| Medicina {n} | 1 |" for n in range(rows)],
        "",
        "| CENA | CENA |",
        "| --- | --- |",
        "| Medicina Z | ½ |",
    ]
    parsed = split_document(lines)
    assert [(s.slot, s.lines, s.table_rows) for s in parsed.sections] == [
        ("desayuno", 0, rows),
        ("cena", 0, 1),
    ]
    sections = [
        {"header": x.header, "slot": x.slot, "lines": x.lines, "table_rows": x.table_rows}
        for x in parsed.sections
    ]
    review = review_list(sections, {"items": [], "avoid": []}, date(2026, 10, 9))
    assert review.warnings == [
        f"table_section_manual_check section=1 rows={rows}",
        "table_section_manual_check section=2 rows=1",
    ]
    # A normal table (plain header row, then separator) counts data rows only.
    parsed = split_document(["CENA", "| Hora | Dosis |", "| --- | --- |", "| 6:00 | 1 |"])
    assert (parsed.sections[0].lines, parsed.sections[0].table_rows) == (0, 1)


def test_evitar_table_rows_count_as_avoid_entries():
    sections = [{"header": "EVITAR", "slot": "evitar", "lines": 1, "table_rows": 2}]
    review = review_list(sections, {"items": [], "avoid": [{"substance": "X"}]}, date(2026, 1, 1))
    assert review.warnings == ["possible_missing_avoid lines=3 avoid=1"]
    avoid = [{"substance": s} for s in "XYZ"]
    assert review_list(sections, {"items": [], "avoid": avoid}, date(2026, 1, 1)).warnings == []


def test_review_completeness():
    sections = [
        {"header": "DESAYUNO", "slot": "desayuno", "lines": 3},
        {"header": "CENA", "slot": "cena", "lines": 1},
        {"header": "EVITAR", "slot": "evitar", "lines": 2},
    ]
    items = [
        {"section": 1, "header": "DESAYUNO", "slot": "desayuno"},
        {"section": 1, "header": "DESAYUNO", "slot": "desayuno"},
        {"section": 2, "header": "CENA", "slot": "cena"},
        {"section": None, "header": None, "slot": "otros"},  # slot set by hand
    ]
    data = {"items": items, "avoid": [{"substance": "X"}]}
    review = review_list(sections, data, date(2026, 10, 9))
    assert review.blocking == []
    assert review.warnings == [
        "item 4: no_section",
        "possible_missing_items section=1 lines=3 items=2",
        "possible_missing_avoid lines=2 avoid=1",
    ]
    items.append({"section": 1, "header": "DESAYUNO", "slot": "desayuno"})
    data["avoid"].append({"substance": "Y"})
    assert review_list(sections, data, date(2026, 10, 9)).warnings == ["item 4: no_section"]
    # A bad_* code stays a REVISAR warning until the operator deletes it;
    # numbering counts every entry, like the load's messages.
    items[0]["warnings"] = ["bad_days", "quantity_too_long"]
    items.insert(0, "not an object")
    sections[1]["table_rows"] = 3
    assert review_list(sections, data, date(2026, 10, 9)).warnings == [
        "item 2: bad_days (fix the field, then delete the code)",
        "item 5: no_section",
        "table_section_manual_check section=2 rows=3",
    ]


def test_extraction_errors(docx_path):
    parsed = split_document(docx_to_markdown_lines(docx_path))
    usage = extract.Usage()
    for stop, payload, code in [
        ("refusal", {}, "refusal"),
        ("max_tokens", {}, "truncated"),
        ("end_turn", ["not", "an", "object"], "invalid_json"),
    ]:
        client = SimpleNamespace(beta=SimpleNamespace(messages=FakeMessages(payload, stop)))
        assert extract.extract_document(client, "m", parsed, usage).error == code


# ── CLI and load ─────────────────────────────────────────────────────────────


@pytest.fixture
def cli(monkeypatch, test_engine):
    from sqlalchemy.orm import sessionmaker

    from app.core.config import Settings

    settings = Settings(anthropic_api_key="test-key", extraction_model="claude-opus-5-5")
    monkeypatch.setattr(import_meds, "get_settings", lambda: settings)
    monkeypatch.setattr(import_meds, "SessionLocal", sessionmaker(bind=test_engine))
    messages = FakeMessages(model_output())
    client = SimpleNamespace(beta=SimpleNamespace(messages=messages))
    monkeypatch.setattr(import_meds, "make_client", lambda key: client)
    return SimpleNamespace(settings=settings, messages=messages)


def _dry_run(cli, docx_path, staging, monkeypatch, capsys):
    def no_db():
        raise AssertionError("a dry run must not open a database session")

    monkeypatch.setattr(import_meds, "SessionLocal", no_db)
    args = [str(docx_path), "--patient", "papa", "--dry-run", "--staging", str(staging)]
    assert import_meds.main(args) == 0
    return capsys.readouterr()


def test_dry_run_writes_staging_only(cli, docx_path, tmp_path, monkeypatch, capsys):
    staging_path = tmp_path / "s.json"
    out = _dry_run(cli, docx_path, staging_path, monkeypatch, capsys)
    assert len(cli.messages.calls) == 1
    assert stat.S_IMODE(staging_path.stat().st_mode) == 0o600
    staging = read_staging(staging_path)
    assert staging["patient"] == "papa"
    assert len(staging["list"]["items"]) == 7
    assert staging["list"]["effective_date"] == "2026-09-30"
    assert "PRUEBA SINTÉTICO" not in staging_path.read_text()
    assert "6 sections" in out.out
    assert "3 header/preamble lines skipped, 1 repeated header lines dropped" in out.out
    assert "extracted: 7 items (" in out.out
    assert "estimated cost: US$" in out.out
    # The model's label for item 7 is no detected header: flagged for review.
    assert "REVISAR (2; the load refuses until fixed):" in out.out
    assert "  REVISAR item 7: unknown_header" in out.out
    # The EVITAR section's table rows count as EVITAR entries (no manual check).
    assert "  REVISAR possible_missing_avoid lines=3 avoid=1" in out.out
    assert "table_section_manual_check" not in out.out
    assert "PHI" in out.err or "HEALTH" in out.err
    # Counts and codes only: no medication names on stdout/stderr.
    for name in ("Medicina", "Suplemento", "Sustancia"):
        assert name not in out.out and name not in out.err
    # A second dry run reuses the staging file: no new call.
    _dry_run(cli, docx_path, staging_path, monkeypatch, capsys)
    assert len(cli.messages.calls) == 1


def test_default_staging_path_and_missing_key(cli, docx_path, tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".git").mkdir()
    cli.settings.anthropic_api_key = ""
    assert import_meds.main([str(docx_path), "--patient", "mama", "--dry-run"]) == 1
    assert "ANTHROPIC_API_KEY" in capsys.readouterr().err
    cli.settings.anthropic_api_key = "test-key"
    assert import_meds.main([str(docx_path), "--patient", "mama", "--dry-run"]) == 0
    assert (tmp_path / "private" / "meds-mama.json").exists()
    monkeypatch.setattr(import_meds, "_in_checkout", lambda cwd: False)
    assert import_meds.main([str(docx_path), "--patient", "mama", "--dry-run"]) == 1
    assert "--staging is required" in capsys.readouterr().err


def test_cli_aborts_on_id_line_before_any_call(cli, tmp_path, capsys):
    path = build_docx(tmp_path / "leak.docx", extra=["CIP 0-000-0002"])
    staging = tmp_path / "s.json"
    args = [str(path), "--patient", "papa", "--dry-run", "--staging", str(staging)]
    assert import_meds.main(args) == 1
    assert cli.messages.calls == []
    assert not staging.exists()
    assert "0-000-0002" not in capsys.readouterr().err


def test_cli_aborts_on_unrecognized_preamble_before_any_call(cli, tmp_path, capsys):
    doc = Document()
    _para(doc, ("LISTA DE MEDICAMENTOS", B))
    _para(doc, (FAKE_PATIENT_HEADER, B))
    _para(doc, "DESDE 30 SEP 26")
    _para(doc, "Medicina A 250 mg 1 tableta")  # no header above it
    _para(doc, ("ENCABEZADO RARO", B))
    _para(doc, "Medicina B 1 tableta")
    _para(doc, ("CENA", B))
    _para(doc, "Medicina C 1 tableta")
    path = tmp_path / "preamble.docx"
    doc.save(path)
    staging = tmp_path / "s.json"
    args = [str(path), "--patient", "papa", "--dry-run", "--staging", str(staging)]
    assert import_meds.main([*args, "--accept-warnings"]) == 1
    assert cli.messages.calls == []
    assert not staging.exists()
    out = capsys.readouterr()
    assert "REVISAR preamble_unrecognized line=4 [blocking]" in out.out
    assert "REVISAR preamble_unrecognized line=5 [blocking]" in out.out
    assert "REVISAR preamble_unrecognized line=6 [blocking]" in out.out
    assert "--accept-warnings does not apply" in out.err
    assert "Medicina" not in out.out + out.err and "PRUEBA" not in out.out + out.err


def test_qualified_header_goes_into_instructions():
    parsed = split_document(
        ["DESDE 30 SEP 26", "DESPUÉS DEL DESAYUNO", "Medicina A", "Medicina B", "Medicina C"]
    )
    output = {
        "items": [
            _item("Medicina A", 1, "DESPUÉS DEL DESAYUNO"),
            _item("Medicina B", 1, "DESPUÉS DEL DESAYUNO", instructions="con agua"),
            _item("Medicina C", 1, "DESPUÉS DEL DESAYUNO", instructions="Después del desayuno"),
        ],
        "avoid": [],
        "list_note": None,
    }
    data, _ = extract.normalize(output, parsed)
    assert [(i["slot"], i["instructions"]) for i in data["items"]] == [
        ("desayuno", "Después del desayuno"),
        ("desayuno", "Después del desayuno; con agua"),
        ("desayuno", "Después del desayuno"),  # already said: not repeated
    ]


def test_load_creates_version_one(
    cli, test_db, test_engine, patients, member, docx_path, tmp_path, monkeypatch, capsys
):
    test_db.add(Doctor(name="Dra. Ana Prueba", aliases=[]))
    test_db.commit()
    staging_path = tmp_path / "s.json"
    _dry_run(cli, docx_path, staging_path, monkeypatch, capsys)
    from sqlalchemy.orm import sessionmaker

    monkeypatch.setattr(import_meds, "SessionLocal", sessionmaker(bind=test_engine))
    by = ["--by", member.email]
    load = ["--patient", "papa", "--from-staging", str(staging_path)]

    # The review warning (item 7: unknown_header) refuses the load...
    assert import_meds.main([*load, *by]) == 2
    out = capsys.readouterr().out
    assert "REVISAR item 7: unknown_header" in out and "--accept-warnings" in out
    assert test_db.scalar(select(MedListVersion.id)) is None
    accept = ["--accept-warnings"]

    # Preview: validated and saved, then rolled back.
    assert import_meds.main([*load, "--dry-run", *accept]) == 0
    out = capsys.readouterr().out
    assert "would create version 1 with 7 items" in out
    assert "accepted (--accept-warnings)" in out
    assert test_db.scalar(select(MedListVersion.id)) is None

    assert import_meds.main([*load, *accept]) == 1  # --by is required to save
    assert import_meds.main([str(docx_path), *load, *by, *accept]) == 0
    out = capsys.readouterr().out
    assert "saved: version 1 with 7 items" in out
    assert "prescribers not matched to one doctor: 1" in out

    papa = patients[0]
    version = test_db.scalar(select(MedListVersion))
    assert (version.patient_id, version.version_no) == (papa.id, 1)
    assert version.changed_by == member.id
    assert version.change_summary == CHANGE_SUMMARY
    assert version.effective_date == date(2026, 9, 30)
    assert version.list_note.startswith("| Hora | Medicina A |")
    assert version.snapshot["effective_date"] == "2026-09-30"
    assert version.snapshot["list_note"] == version.list_note
    assert len(version.snapshot["items"]) == 7

    items = {i.name: i for i in meds.current_items(test_db, papa.id)}
    assert [i.name for i in meds.current_items(test_db, papa.id)][:3] == [
        "Medicina A",
        "Medicina B",
        "Medicina C",
    ]
    assert items["Medicina A"].quantity == "½" and items["Medicina A"].strength == "250 mg"
    assert items["Medicina B"].days == [1, 3, 5]
    assert items["Medicina B"].prescriber_id is not None
    assert items["Medicina C"].instructions == "Indicado por: Dr. Desconocido"
    assert items["Medicina C"].schedule_note.startswith("vie 25 sep")
    assert (items["Medicina D"].slot, items["Medicina D"].time) == ("noche", time(19, 30))
    assert items["Medicina E"].slot == "segun_necesidad"
    assert items["Medicina C"].sort_order == 1
    assert items["Medicina A"].created_by == member.id
    assert [r.substance for r in test_db.scalars(select(MedAvoid))] == ["Sustancia X"]

    # Idempotency: a second load (or a full run) is refused, nothing sent.
    assert import_meds.main([*load, *by, *accept]) == 1
    assert "already has a medication list (version 1)" in capsys.readouterr().err
    calls = len(cli.messages.calls)
    monkeypatch.setattr(import_meds, "make_client", lambda key: pytest.fail("no call"))
    other = tmp_path / "other.json"
    assert (
        import_meds.main([str(docx_path), "--patient", "papa", "--staging", str(other), *by]) == 1
    )
    assert len(cli.messages.calls) == calls and not other.exists()
    assert test_db.scalar(select(MedListVersion.version_no).order_by(MedListVersion.id.desc())) == 1


def test_mismatches(cli, test_db, patients, member, docx_path, tmp_path, monkeypatch, capsys):
    staging_path = tmp_path / "s.json"
    _dry_run(cli, docx_path, staging_path, monkeypatch, capsys)
    by = ["--by", member.email]
    assert import_meds.main(["--patient", "mama", "--from-staging", str(staging_path), *by]) == 1
    assert "not 'mama'" in capsys.readouterr().err
    other = tmp_path / "other.docx"
    Document().save(other)
    args = [str(other), "--patient", "papa", "--from-staging", str(staging_path), *by]
    assert import_meds.main(args) == 1
    assert "does not match" in capsys.readouterr().err
    bitacora_like = tmp_path / "b.json"
    bitacora_like.write_text(json.dumps({"version": 1, "entries": []}))
    assert import_meds.main(["--patient", "papa", "--from-staging", str(bitacora_like), *by]) == 1
    # Another document's staging at the --staging path is never overwritten.
    path2 = build_docx(tmp_path / "otra.docx", extra=["Medicina Z 1"])
    args = [str(path2), "--patient", "papa", "--dry-run", "--staging", str(staging_path)]
    assert import_meds.main(args) == 1
    assert "delete it or choose another" in capsys.readouterr().err


def _staging(items, avoid=(), **list_fields):
    return {
        "kind": "meds",
        "version": 1,
        "source_sha256": "a" * 64,
        "patient": "papa",
        "sections": [{"header": "DESAYUNO", "slot": "desayuno", "lines": 1}],
        "list": {
            "effective_date": None,
            "list_note": None,
            "items": list(items),
            "avoid": list(avoid),
            **list_fields,
        },
        "error": None,
    }


def test_invalid_staging_saves_nothing(test_db, patients):
    staging = _staging(
        [
            {"section": 1, "name": "Medicina A", "slot": "desayuno", "quantity": "1"},
            {"section": 1, "name": "Medicina B", "slot": "merienda"},  # not a slot
            {"section": 1, "name": "", "slot": "cena"},
            {"section": 1, "name": "Medicina C", "slot": "desayuno", "quantity": "x" * 60},
            {"section": 1, "name": "Medicina D", "slot": None},  # from EVITAR / no section
            {
                "section": 1,
                "name": "Medicina E",
                "slot": "desayuno",
                "start_date": "2026-10-05",
                "end_date": "2026-10-01",
            },
        ],
        effective_date="2026-13-01",
    )
    report = load_staging(test_db, staging, patient=patients[0], user_id=None, accept_warnings=True)
    assert not report.saved
    assert "item 2: slot literal_error" in report.failures
    assert "item 3: name string_too_short" in report.failures
    assert "item 4: quantity string_too_long" in report.failures  # never cut
    assert any(f.startswith("item 5: needs_review") for f in report.failures)
    assert "item 6: - value_error" in report.failures  # end before start
    assert "list: effective_date invalid" in report.failures
    assert test_db.scalar(select(MedItem.id)) is None
    empty = load_staging(test_db, _staging([]), patient=patients[0], user_id=None)
    assert empty.failures == ["REVISAR possible_missing_items section=1 lines=1 items=0"]
    empty = load_staging(
        test_db, _staging([]), patient=patients[0], user_id=None, accept_warnings=True
    )
    assert empty.failures == ["list: no items"]


def test_past_end_date_is_reported(test_db, patients):
    staging = _staging(
        [
            {"section": 1, "name": "Medicina A", "slot": "desayuno"},
            {"section": 1, "name": "Medicina B", "slot": "desayuno", "end_date": "2026-01-01"},
        ]
    )
    report = load_staging(test_db, staging, patient=patients[0], user_id=None)
    assert not report.saved
    assert report.failures == ["REVISAR item 2: ends_in_past (would not be on the current list)"]
    report = load_staging(test_db, staging, patient=patients[0], user_id=None, accept_warnings=True)
    assert report.saved and report.failures == []
    assert [i.name for i in meds.current_items(test_db, patients[0].id)] == ["Medicina A"]


def test_later_save_carries_the_list_note_and_desde(auth_client, test_db, patients):
    staging = _staging(
        [
            {
                "section": 1,
                "name": "Medicina A",
                "slot": "desayuno",
                "quantity": "1",
                "header": "[S1] DESAYUNO",
            }
        ],
        list_note="Nota sintética",
        effective_date="2026-09-30",
    )
    papa = patients[0]
    assert load_staging(test_db, staging, patient=papa, user_id=None).saved
    body = auth_client.get(f"/api/v1/patients/{papa.id}/meds").json()
    assert (body["list_note"], body["effective_date"]) == ("Nota sintética", "2026-09-30")
    item = body["slots"][1]["items"][0]
    resp = auth_client.put(
        f"/api/v1/patients/{papa.id}/meds",
        json={
            "expected_version": 1,
            "items": [
                {"id": item["id"], "name": "Medicina A", "slot": "desayuno", "quantity": "½"}
            ],
            "avoid": [],
        },
    )
    assert resp.status_code == 200, resp.text
    # #135: an editor save without the list fields keeps both.
    assert (resp.json()["list_note"], resp.json()["effective_date"]) == (
        "Nota sintética",
        "2026-09-30",
    )
