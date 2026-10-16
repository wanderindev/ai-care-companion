"""Chat file tools (generate_pdf / generate_xlsx), synthetic data only.

The PDF tests render through the real child process (WeasyPrint); storage is
moto (``storage`` fixture). The turn-level tests (SSE ``file`` event, linking
at persist) are in ``test_chat.py`` next to the fake Anthropic client.
"""

import io
import json
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.models.attachment import Attachment
from app.services.appointments import today_local
from app.services.chat import files, tools
from app.services.pdf.chat_document import (
    ChatDocument,
    DocSection,
    DocTable,
    render_chat_document,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return " ".join(" ".join(p.extract_text() or "" for p in reader.pages).split())


def pdf_args(**kw):
    return {
        "title": "Resumen de citas",
        "sections": [{"heading": "Próximas citas", "markdown": "- Control el 12/03/2025"}],
        **kw,
    }


def xlsx_args(**kw):
    return {
        "title": "Laboratorios",
        "sheets": [
            {"name": "Glucosa", "columns": ["Fecha", "Valor"], "rows": [["2025-03-12", 98]]}
        ],
        **kw,
    }


def spec_error(parse, args) -> str:
    with pytest.raises(files.FileSpecError) as info:
        parse(args)
    return str(info.value)


# ── Input validation and caps ───────────────────────────────────────────────


def test_pdf_input_is_validated():
    assert "title" in spec_error(files.parse_pdf, pdf_args(title="  "))
    assert "sections" in spec_error(files.parse_pdf, pdf_args(sections=[]))
    assert "sections" in spec_error(files.parse_pdf, pdf_args(sections="texto"))
    assert "markdown" in spec_error(files.parse_pdf, pdf_args(sections=[{"heading": "Solo"}]))
    assert "desconocidos" in spec_error(files.parse_pdf, pdf_args(sections=[{"html": "<b>"}]))
    bad_cell = {"table": {"columns": ["A"], "rows": [[{"x": 1}]]}}
    assert "celdas" in spec_error(files.parse_pdf, pdf_args(sections=[bad_cell]))
    wide_row = {"table": {"columns": ["A"], "rows": [["1", "2"]]}}
    assert "más celdas" in spec_error(files.parse_pdf, pdf_args(sections=[wide_row]))
    assert "title" in spec_error(files.parse_pdf, pdf_args(title="x" * 201))


def test_pdf_caps_ask_to_shorten():
    many = [{"markdown": "x"}] * (files.MAX_SECTIONS + 1)
    assert "máximo 50" in spec_error(files.parse_pdf, pdf_args(sections=many))
    long_md = [{"markdown": "x" * (files.MAX_MARKDOWN_CHARS + 1)}]
    assert "acórtalo" in spec_error(files.parse_pdf, pdf_args(sections=long_md))
    total = [{"markdown": "x" * files.MAX_MARKDOWN_CHARS}] * 4
    assert "resume" in spec_error(files.parse_pdf, pdf_args(sections=total))
    rows = {"table": {"columns": ["A"], "rows": [["1"]] * (files.MAX_PDF_ROWS + 1)}}
    assert "filas" in spec_error(files.parse_pdf, pdf_args(sections=[rows]))
    cols = {"table": {"columns": ["c"] * (files.MAX_PDF_COLUMNS + 1), "rows": []}}
    assert "columnas" in spec_error(files.parse_pdf, pdf_args(sections=[cols]))
    big = {"table": {"columns": ["c"] * 20, "rows": [["v"] * 20] * 199}}
    assert "celdas" in spec_error(files.parse_pdf, pdf_args(sections=[big, big]))
    cell = {"table": {"columns": ["A"], "rows": [["x" * (files.MAX_CELL_CHARS + 1)]]}}
    assert "acórtala" in spec_error(files.parse_pdf, pdf_args(sections=[cell]))


def test_pdf_cells_become_text_and_short_rows_are_padded():
    spec = files.parse_pdf(
        pdf_args(
            sections=[
                {"table": {"columns": ["A", "B", "C"], "rows": [[1.0, True, None], [2.5]]}},
            ]
        )
    )
    assert spec.sections[0].table.rows == [["1", "Sí", ""], ["2.5", "", ""]]


def test_xlsx_caps():
    many = [{"name": f"H{i}", "columns": ["A"], "rows": []} for i in range(files.MAX_SHEETS + 1)]
    assert "hojas" in spec_error(files.parse_xlsx, xlsx_args(sheets=many))
    wide = [{"name": "A", "columns": ["c"] * (files.MAX_XLSX_COLUMNS + 1), "rows": []}]
    assert "columnas" in spec_error(files.parse_xlsx, xlsx_args(sheets=wide))
    big = [{"name": "A", "columns": ["a", "b"], "rows": [[1, 2]] * files.MAX_CELLS}]
    assert "5000 celdas" in spec_error(files.parse_xlsx, xlsx_args(sheets=big))
    assert "sheets" in spec_error(files.parse_xlsx, xlsx_args(sheets=[]))


def test_sheet_names_follow_excel_rules():
    used: set[str] = set()
    assert files.sheet_name("Laboratorio [2025]: a/b?*\\", 1, used) == "Laboratorio 2025 a b"
    assert files.sheet_name("'Citas'", 2, used) == "Citas"
    assert files.sheet_name("citas", 3, used) == "citas (2)"
    assert files.sheet_name("", 4, used) == "Hoja 4"
    assert files.sheet_name("History", 5, used) == "Historial"
    long = files.sheet_name("Resultados de laboratorio de toda la familia", 6, used)
    assert long == "Resultados de laboratorio de to" and len(long) == 31
    again = files.sheet_name("Resultados de laboratorio de toda la familia", 7, used)
    assert again.endswith(" (2)") and len(again) <= 31


def test_output_filename():
    day = date(2026, 10, 9)
    assert files.output_filename(None, "Resumen de citas — Papá", "pdf", day) == (
        "resumen-de-citas-papa-2026-10-09.pdf"
    )
    assert files.output_filename("Labs Mamá.XLSX", "x", "xlsx", day) == "labs-mama-2026-10-09.xlsx"
    assert files.output_filename("labs-2025-01-01", "x", "xlsx", day) == "labs-2026-10-09.xlsx"
    assert files.output_filename("../../etc/passwd", "x", "pdf", day) == "etc-passwd-2026-10-09.pdf"
    assert files.output_filename("¿¡!?", "★", "pdf", day) == "documento-2026-10-09.pdf"


def test_quote_prefix_rule():
    for value in ('=HYPERLINK("http://x")', "+1+2", "-2+3", "@SUM(A1)", "\tx"):
        assert files.needs_quote_prefix(value)
    for value in ("-5", "+3.2", "12,5", "Normal", ""):
        assert not files.needs_quote_prefix(value)


# ── XLSX ────────────────────────────────────────────────────────────────────


def test_xlsx_round_trip_is_safe_and_formatted():
    spec = files.parse_xlsx(
        {
            "title": "Laboratorios",
            "sheets": [
                {
                    "name": "Glucosa",
                    "columns": ["Fecha", "Prueba", "Valor", "En rango", "Nota"],
                    "rows": [
                        ["2025-03-12", "Glucosa", 98, True, '=HYPERLINK("http://x","y")'],
                        ["2025-13-40", "Creatinina", 0.9, False, "@cmd"],
                        ["2025-05-02", "x" * 500, -1, None, "-5\x07"],
                    ],
                },
                {"name": "Vacía", "columns": ["A"], "rows": []},
            ],
        }
    )
    data = files.build_xlsx(spec)
    sheet_xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml  # no formula anywhere

    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Glucosa", "Vacía"]
    assert wb.properties.title == "Laboratorios"
    ws = wb["Glucosa"]
    assert [c.value for c in ws[1]] == ["Fecha", "Prueba", "Valor", "En rango", "Nota"]
    assert all(c.font.bold for c in ws[1])
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:E4"
    assert ws["A2"].value == datetime(2025, 3, 12) and ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["A3"].value == "2025-13-40"  # not a real date: stays text
    assert ws["C2"].value == 98 and ws["C3"].value == 0.9 and ws["D2"].value is True
    formula = ws["E2"]
    assert formula.value == '=HYPERLINK("http://x","y")'
    assert formula.data_type == "s" and formula.quotePrefix
    assert ws["E3"].value == "@cmd" and ws["E3"].quotePrefix
    assert ws["E4"].value == "-5" and not ws["E4"].quotePrefix  # control char dropped
    assert ws.column_dimensions["B"].width == files.MAX_COLUMN_WIDTH
    assert ws.column_dimensions["D"].width == len("En rango") + 2
    assert wb["Vacía"].column_dimensions["A"].width == files.MIN_COLUMN_WIDTH
    assert wb["Vacía"].auto_filter.ref is None


# ── PDF rendering (real child process) ──────────────────────────────────────


def test_render_chat_document_pdf():
    doc = ChatDocument(
        title="Resumen para la consulta",
        patient_name="Papá",
        sections=[
            DocSection(
                heading="Indicaciones",
                markdown="**Medicina A** dos veces al día.\n\n<b>sin html</b>\n\n![x](http://x/y.png)",
            ),
            DocSection(
                heading="Glucosa",
                table=DocTable(columns=["Fecha", "Valor"], rows=[["12/03/2025", "98 mg/dL"]]),
            ),
        ],
    )
    content = render_chat_document(doc, generated_at=datetime(2026, 10, 9, 15, 30))
    assert content.startswith(b"%PDF-")
    text = pdf_text(content)
    assert "Resumen para la consulta" in text and "Papá" in text
    assert "Generado el 9 de octubre de 2026" in text
    assert "Indicaciones" in text and "Medicina A dos veces al día." in text
    assert "<b>sin html</b>" in text  # raw HTML stays text
    assert "12/03/2025" in text and "98 mg/dL" in text


# ── The tools (database + moto storage) ─────────────────────────────────────


def run(db, name, args, storage=None, user_id=None):
    ctx = tools.ToolContext(db=db, user_id=user_id, storage=(lambda: storage) if storage else None)
    result = tools.run_tool(ctx, "toolu_f", name, args)
    content = result.block["content"]
    return ctx, result, (json.loads(content) if result.ok else content)


def test_generate_pdf_stores_an_unlinked_attachment(test_db, patients, storage, member):
    papa = patients[0]
    papa.full_name = "Nombre Completo Sintético"
    papa.id_number = "0-000-0000"
    test_db.commit()
    ctx, result, out = run(
        test_db,
        "generate_pdf",
        pdf_args(patient="papa", filename="Resumen citas"),
        storage,
        user_id=member.id,
    )
    assert result.ok, out
    name = f"resumen-citas-{today_local().isoformat()}.pdf"
    assert out["filename"] == name and out["file_id"] == ctx.attachment_ids[0]
    assert out["size"] > 1000 and "descarga" in out["note"]
    row = test_db.get(Attachment, out["file_id"])
    assert (row.owner_type, row.owner_id) == (None, None)
    assert (row.kind, row.mime, row.filename) == ("documents", "application/pdf", name)
    assert row.uploaded_by == member.id
    assert ctx.files == [
        {"id": row.id, "filename": name, "mime": "application/pdf", "size": row.size}
    ]
    stored = b"".join(storage.iter_object(row.storage_key))
    text = pdf_text(stored)
    assert "Papá" in text and "Próximas citas" in text
    # Display name only: never the full name or the CIP.
    assert "Sintético" not in text and "0-000-0000" not in text


def test_generate_xlsx_stores_the_workbook(test_db, patients, storage):
    ctx, result, out = run(test_db, "generate_xlsx", xlsx_args(), storage)
    assert result.ok, out
    row = test_db.get(Attachment, out["file_id"])
    assert row.mime == XLSX_MIME and row.kind == "documents"
    assert row.storage_key.endswith(".xlsx")
    assert out["filename"] == f"laboratorios-{today_local().isoformat()}.xlsx"
    wb = load_workbook(io.BytesIO(b"".join(storage.iter_object(row.storage_key))))
    assert wb["Glucosa"]["B2"].value == 98


def test_file_tool_errors_store_nothing(test_db, patients, storage):
    _, result, out = run(test_db, "generate_pdf", pdf_args(sections=[]), storage)
    assert not result.ok and "sections" in out
    _, result, out = run(test_db, "generate_pdf", pdf_args(patient="tio"), storage)
    assert not result.ok and "Paciente desconocido" in out
    _, result, out = run(test_db, "generate_xlsx", xlsx_args(extra=1), storage)
    assert not result.ok and "extra" in out
    _, result, out = run(test_db, "generate_xlsx", xlsx_args(), None)
    assert not result.ok and out == tools.STORAGE_UNAVAILABLE
    assert test_db.scalars(select(Attachment)).all() == []


def test_storage_failure_is_an_error_result_without_details(test_db, patients, caplog):
    class Broken:
        def exists(self, key):
            raise RuntimeError("Laboratorios secretos")

    with caplog.at_level("ERROR"):
        _, result, out = run(test_db, "generate_xlsx", xlsx_args(), Broken())
    assert not result.ok and out == tools.STORE_FAILED
    assert "RuntimeError" in caplog.text and "secretos" not in caplog.text
    assert test_db.scalars(select(Attachment)).all() == []


def test_no_connection_is_held_while_rendering(
    test_db, patients, storage, test_engine, monkeypatch
):
    seen = []

    def fake_render(doc):
        seen.append(test_engine.pool.checkedout())
        return b"%PDF-1.7 synthetic"

    monkeypatch.setattr(tools, "render_chat_document", fake_render)
    test_db.close()  # the fixture's own session holds nothing either
    _, result, out = run(test_db, "generate_pdf", pdf_args(patient="mama"), storage)
    assert result.ok, out
    assert seen == [0]
