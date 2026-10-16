"""The chat's file tools (#33): ``generate_pdf`` and ``generate_xlsx``.

This module validates the tool input (shape and size caps), builds the
workbook and names the file; :mod:`app.services.chat.tools` resolves the
patient, renders and stores. Everything here is pure (no database).

Size caps keep a turn's memory and time bounded; over a cap the tool answers
with an error that asks Claude to shorten (:class:`FileSpecError`):

- PDF: ≤ ``MAX_SECTIONS`` sections, markdown ≤ ``MAX_MARKDOWN_CHARS`` per
  section and ``MAX_TOTAL_MARKDOWN_CHARS`` in all, tables ≤ ``MAX_PDF_ROWS``
  rows × ``MAX_PDF_COLUMNS`` columns, ≤ ``MAX_CELLS`` table cells in all.
- XLSX: ≤ ``MAX_SHEETS`` sheets of ≤ ``MAX_XLSX_COLUMNS`` columns,
  ≤ ``MAX_CELLS`` cells in the workbook (headers included).
- Text cells ≤ ``MAX_CELL_CHARS``; titles and headings ≤ ``MAX_TITLE_CHARS``.

Spreadsheet safety: the content comes from the model and the data, so no
cell is ever a formula. Strings are written as text cells, and a string that
Excel or a CSV re-import would read as a formula (it starts with ``=``,
``+``, ``-``, ``@``, a tab or a CR and is not a plain number) also gets
Excel's «quote prefix» (shown as text, kept as text when edited). Control
characters XML cannot carry are dropped. Strings that are exactly an ISO
date (``AAAA-MM-DD``) become date cells shown as ``dd/mm/aaaa``; numbers
and booleans stay numbers and booleans.
"""

import io
import math
import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from typing import Any

from app.services.pdf.chat_document import DocSection, DocTable

MAX_TITLE_CHARS = 200
MAX_SECTIONS = 50
MAX_MARKDOWN_CHARS = 20_000
MAX_TOTAL_MARKDOWN_CHARS = 60_000
MAX_PDF_ROWS = 200
MAX_PDF_COLUMNS = 20
MAX_SHEETS = 10
MAX_XLSX_COLUMNS = 30
MAX_CELLS = 5_000
MAX_CELL_CHARS = 2_000
MAX_FILENAME_CHARS = 60

SHEET_NAME_MAX = 31
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
# XML 1.0 cannot carry these (openpyxl refuses them).
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_PLAIN_NUMBER = re.compile(r"^[+-]?\d+([.,]\d+)?$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_TRAILING_DATE = re.compile(r"-\d{4}-\d{2}-\d{2}$")

MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
DATE_FORMAT = "DD/MM/YYYY"


class FileSpecError(ValueError):
    """Invalid or oversized input; the message goes back to Claude."""


# ── Common ──────────────────────────────────────────────────────────────────


def _text(value: Any, name: str, *, required: bool = False, max_len: int = MAX_TITLE_CHARS):
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise FileSpecError(f'Falta "{name}".')
        return None
    if not isinstance(value, str):
        raise FileSpecError(f'"{name}" debe ser texto.')
    value = _ILLEGAL_XML.sub("", value).strip()
    if len(value) > max_len:
        raise FileSpecError(f'"{name}" es demasiado largo (máximo {max_len} caracteres): acórtalo.')
    return value


def _list(value: Any, name: str) -> list[Any]:
    if not isinstance(value, list):
        raise FileSpecError(f'"{name}" debe ser una lista.')
    return value


def _object(value: Any, name: str, allowed: set[str]) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise FileSpecError(f'"{name}" debe ser un objeto.')
    unknown = set(value) - allowed
    if unknown:
        raise FileSpecError(f'Campos desconocidos en "{name}": {", ".join(sorted(unknown))}.')
    return value


def _check_cell(value: Any, where: str) -> Any:
    if value is None or isinstance(value, bool | int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, str):
        value = _ILLEGAL_XML.sub("", value)
        if len(value) > MAX_CELL_CHARS:
            raise FileSpecError(
                f"Una celda de {where} supera {MAX_CELL_CHARS} caracteres: acórtala."
            )
        return value
    raise FileSpecError(f"Las celdas de {where} deben ser texto, número, booleano o null.")


def _table(columns: Any, rows: Any, where: str, *, max_rows: int | None, max_cols: int):
    columns = _list(columns, f"columns de {where}")
    if not columns:
        raise FileSpecError(f'"columns" de {where} no puede estar vacío.')
    if len(columns) > max_cols:
        raise FileSpecError(
            f"{where} tiene {len(columns)} columnas (máximo {max_cols}): reduce las columnas."
        )
    headers = [_text(c, f"columns de {where}") or "" for c in columns]
    rows = _list(rows, f"rows de {where}")
    if max_rows is not None and len(rows) > max_rows:
        raise FileSpecError(
            f"{where} tiene {len(rows)} filas (máximo {max_rows}): resume o divide la tabla."
        )
    out: list[list[Any]] = []
    for row in rows:
        row = _list(row, f"una fila de {where}")
        if len(row) > len(headers):
            raise FileSpecError(f"Una fila de {where} tiene más celdas que columnas.")
        cells = [_check_cell(v, where) for v in row]
        out.append(cells + [None] * (len(headers) - len(cells)))
    return headers, out


def slugify(text: str) -> str:
    folded = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", folded.lower()).strip("-")
    return slug[:MAX_FILENAME_CHARS].strip("-")


def output_filename(requested: str | None, title: str, ext: str, today: date) -> str:
    """``<slug of the requested name or the title>-<yyyy-mm-dd>.<ext>``."""
    base = requested or ""
    base = re.sub(rf"\.{ext}$", "", base.strip(), flags=re.IGNORECASE)
    slug = _TRAILING_DATE.sub("", slugify(base)) or _TRAILING_DATE.sub("", slugify(title))
    return f"{slug or 'documento'}-{today.isoformat()}.{ext}"


# ── PDF ─────────────────────────────────────────────────────────────────────


@dataclass
class PdfSpec:
    title: str
    patient: Any
    sections: list[DocSection]
    filename: str | None


def _pdf_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_pdf(args: dict[str, Any]) -> PdfSpec:
    title = _text(args.get("title"), "title", required=True)
    filename = _text(args.get("filename"), "filename")
    sections_in = _list(args.get("sections"), "sections")
    if not sections_in:
        raise FileSpecError('"sections" no puede estar vacío.')
    if len(sections_in) > MAX_SECTIONS:
        raise FileSpecError(
            f"Hay {len(sections_in)} secciones (máximo {MAX_SECTIONS}): agrúpalas o resume."
        )
    sections: list[DocSection] = []
    total_md = 0
    total_cells = 0
    for i, raw in enumerate(sections_in, start=1):
        where = f"la sección {i}"
        raw = _object(raw, f"sections[{i - 1}]", {"heading", "markdown", "table"})
        heading = _text(raw.get("heading"), f"heading de {where}")
        markdown = _text(raw.get("markdown"), f"markdown de {where}", max_len=MAX_MARKDOWN_CHARS)
        total_md += len(markdown or "")
        table = None
        if raw.get("table") is not None:
            t = _object(raw["table"], f"table de {where}", {"columns", "rows"})
            headers, rows = _table(
                t.get("columns"),
                t.get("rows"),
                f"la tabla de {where}",
                max_rows=MAX_PDF_ROWS,
                max_cols=MAX_PDF_COLUMNS,
            )
            total_cells += len(headers) * (len(rows) + 1)
            table = DocTable(columns=headers, rows=[[_pdf_cell(v) for v in r] for r in rows])
        if not markdown and table is None:
            raise FileSpecError(f'{where.capitalize()} necesita "markdown" o "table".')
        sections.append(DocSection(heading=heading, markdown=markdown, table=table))
    if total_md > MAX_TOTAL_MARKDOWN_CHARS:
        raise FileSpecError(
            f"El texto suma {total_md} caracteres (máximo {MAX_TOTAL_MARKDOWN_CHARS}): resume."
        )
    if total_cells > MAX_CELLS:
        raise FileSpecError(
            f"Las tablas suman {total_cells} celdas (máximo {MAX_CELLS}): resume o divide."
        )
    return PdfSpec(title=title, patient=args.get("patient"), sections=sections, filename=filename)


# ── XLSX ────────────────────────────────────────────────────────────────────


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[Any]]


@dataclass
class XlsxSpec:
    title: str
    sheets: list[Sheet]
    filename: str | None


def sheet_name(raw: str | None, index: int, used: set[str]) -> str:
    """An Excel-valid, unique sheet name: no ``[]:*?/\\``, not starting or
    ending with an apostrophe, ≤ 31 characters, not «History» (reserved)."""
    name = _SHEET_FORBIDDEN.sub(" ", raw or "")
    name = re.sub(r"\s+", " ", name).strip().strip("'").strip()
    name = name[:SHEET_NAME_MAX].strip()
    if not name or name.lower() == "history":
        name = f"Hoja {index}" if not name else "Historial"
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = name[: SHEET_NAME_MAX - len(suffix)].rstrip() + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def parse_xlsx(args: dict[str, Any]) -> XlsxSpec:
    title = _text(args.get("title"), "title", required=True)
    filename = _text(args.get("filename"), "filename")
    sheets_in = _list(args.get("sheets"), "sheets")
    if not sheets_in:
        raise FileSpecError('"sheets" no puede estar vacío.')
    if len(sheets_in) > MAX_SHEETS:
        raise FileSpecError(f"Hay {len(sheets_in)} hojas (máximo {MAX_SHEETS}): agrúpalas.")
    used: set[str] = set()
    sheets: list[Sheet] = []
    total_cells = 0
    for i, raw in enumerate(sheets_in, start=1):
        raw = _object(raw, f"sheets[{i - 1}]", {"name", "columns", "rows"})
        name = sheet_name(_text(raw.get("name"), f"name de la hoja {i}"), i, used)
        headers, rows = _table(
            raw.get("columns"),
            raw.get("rows"),
            f"la hoja {i}",
            max_rows=None,
            max_cols=MAX_XLSX_COLUMNS,
        )
        total_cells += len(headers) * (len(rows) + 1)
        if total_cells > MAX_CELLS:
            raise FileSpecError(
                f"El libro supera {MAX_CELLS} celdas: resume, filtra o divide en varios archivos."
            )
        sheets.append(Sheet(name=name, columns=headers, rows=rows))
    return XlsxSpec(title=title, sheets=sheets, filename=filename)


def needs_quote_prefix(text: str) -> bool:
    return text.startswith(_FORMULA_START) and not _PLAIN_NUMBER.match(text)


def _iso_date(text: str) -> date | None:
    if not _ISO_DATE.match(text):
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _write_cell(cell, value: Any) -> int:
    """Set ``cell`` safely; returns the display width it needs."""
    if value is None:
        return 0
    if isinstance(value, bool | int | float):
        cell.value = value
        return len(str(value))
    as_date = _iso_date(value)
    if as_date is not None:
        cell.value = as_date
        cell.number_format = DATE_FORMAT
        return 10
    cell.value = value
    cell.data_type = "s"  # text, never a formula
    if needs_quote_prefix(value):
        cell.quotePrefix = True
    longest = max((len(line) for line in value.splitlines()), default=0)
    return longest


def build_xlsx(spec: XlsxSpec) -> bytes:
    """The workbook as .xlsx bytes: one sheet per input sheet, bold header
    row, frozen header, autofilter, column widths fitted (bounded)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = spec.title
    wb.properties.creator = "n16chat"
    bold = Font(bold=True)
    for sheet in spec.sheets:
        ws = wb.create_sheet(title=sheet.name)
        widths = [0] * len(sheet.columns)
        for col, header in enumerate(sheet.columns, start=1):
            cell = ws.cell(row=1, column=col)
            widths[col - 1] = _write_cell(cell, header)
            cell.font = bold
        for r, row in enumerate(sheet.rows, start=2):
            for col, value in enumerate(row, start=1):
                width = _write_cell(ws.cell(row=r, column=col), value)
                widths[col - 1] = max(widths[col - 1], width)
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = max(
                MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, width + 2)
            )
        ws.freeze_panes = "A2"
        if sheet.rows:
            ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
